"""
roi_table_dialog.py  (v2 — Pivot-style interactive ROI statistics table)

Layout
------
  Filter bar  : Pool checkboxes + Method checkboxes  (live rebuild)
  Two-level header:
      Row 0 (data area)  — Pool name, colour-coded, spans all method columns
      QHeaderView        — Method names  (click to select whole column)
  Data rows   : ROI name  |  Mean ± Std per (pool, method)  |  n pixels
  Selection   : ExtendedSelection / SelectItems  → cells, rows, columns all work
  Export      : Copy CSV  |  Export Excel (.xlsx)
"""
from __future__ import annotations

import numpy as np
from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QPushButton,
    QTableWidget, QTableWidgetItem, QLabel, QCheckBox,
    QAbstractItemView, QHeaderView, QSizePolicy, QMessageBox,
    QFrame, QWidget, QLayout, QSpinBox, QComboBox, QColorDialog,
    QFileDialog, QLineEdit,
)
from PyQt6.QtCore import Qt, QRect, QSize, QPoint, pyqtSignal, QTimer
from PyQt6.QtGui import QFont, QColor, QBrush, QPalette


class _FlowLayout(QLayout):
    """Left-to-right layout that wraps to a new line when width is exceeded
    (Qt's canonical FlowLayout). Keeps the filter controls on one line when
    there are few, and only wraps when there are many."""

    def __init__(self, parent=None, spacing=10):
        super().__init__(parent)
        self.setSpacing(spacing)
        self._items = []

    def addItem(self, item):        self._items.append(item)
    def count(self):                return len(self._items)
    def itemAt(self, i):            return self._items[i] if 0 <= i < len(self._items) else None
    def takeAt(self, i):            return self._items.pop(i) if 0 <= i < len(self._items) else None
    def expandingDirections(self):  return Qt.Orientation(0)
    def hasHeightForWidth(self):    return True
    def heightForWidth(self, w):    return self._do(QRect(0, 0, w, 0), True)

    def setGeometry(self, rect):
        super().setGeometry(rect)
        self._do(rect, False)

    def sizeHint(self):             return self.minimumSize()

    def minimumSize(self):
        s = QSize()
        for it in self._items:
            s = s.expandedTo(it.minimumSize())
        m = self.contentsMargins()
        return s + QSize(m.left() + m.right(), m.top() + m.bottom())

    def _do(self, rect, test):
        x, y, line_h = rect.x(), rect.y(), 0
        sp = self.spacing()
        for it in self._items:
            w, h = it.sizeHint().width(), it.sizeHint().height()
            if x + w > rect.right() and line_h > 0:
                x = rect.x(); y += line_h + sp; line_h = 0
            if not test:
                it.setGeometry(QRect(QPoint(x, y), it.sizeHint()))
            x += w + sp
            line_h = max(line_h, h)
        return y + line_h - rect.y()


# ─────────────────────────────────────────────────────────────────────────────
# Colour palette — one background per pool group
# ─────────────────────────────────────────────────────────────────────────────
# Pool-header row — light pastel backgrounds, dark text for readability
_POOL_BG = [
    "#cce5ff",   # sky blue       (water)
    "#ccf5e0",   # mint green     (amine / amide)
    "#ffe5cc",   # peach          (NOE)
    "#e8d5f5",   # lavender       (MT)
    "#ffd6e0",   # blush pink     (OH)
    "#ccf5f5",   # aqua           (guanidinium / glucose)
    "#fffacc",   # lemon          (creatine / taurine)
    "#e0e0e0",   # light gray     (Other / PLL)
]
_POOL_TEXT = "#1a1a1a"   # dark text on all light backgrounds (light OS theme)


def _is_dark_theme() -> bool:
    """True if the OS/app is using a dark theme (so text should be light)."""
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance()
    if app is None:
        return True
    return app.palette().color(QPalette.ColorRole.Window).lightness() < 128


def _pool_header_colors(pastel_hex: str, is_dark: bool) -> tuple[str, str]:
    """Return (background, text) for a pool-group header that is readable on
    either OS theme: white text on a darkened pool tint for dark themes, dark
    text on the light pastel for light themes."""
    if is_dark:
        return QColor(pastel_hex).darker(300).name(), "#ffffff"
    return pastel_hex, "#1a1a1a"

# Display-friendly pool names
_POOL_DISPLAY_NAMES: dict[str, str] = {
    'NOE':            'NOE',
    'MT':             'MT',
    'OH':             'OH',
    'poly_l_lysine':  'Poly-L-Lysine',
    'guanidinium':    'Guanidinium',
    'iopamidol_4.2':  'Iopamidol 4.2',
    'iopamidol_5.5':  'Iopamidol 5.5',
    'Other':          'Other',
}

def _display_pool_name(pool: str) -> str:
    """Capitalise pool name for display (preserves all-caps like NOE/MT/OH)."""
    if pool in _POOL_DISPLAY_NAMES:
        return _POOL_DISPLAY_NAMES[pool]
    # Keep special labels verbatim (e.g. "%CEST @4.20ppm", "B0 (Hz) [400 MHz]")
    if any(s in pool for s in ('%', '[')) or 'ppm' in pool.lower() or 'Hz' in pool or pool.startswith('B0'):
        return pool
    return pool.replace('_', ' ').title()


_POOL_BG_LIGHT = _POOL_BG   # kept for any external references


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _roi_stats(roi, map_data: np.ndarray) -> tuple[float, float, int] | None:
    """Return (mean, std, n) of finite pixel values inside roi mask, or None."""
    try:
        msk = roi.mask
        if msk.shape != map_data.shape[:2]:
            from scipy.ndimage import zoom
            zy = map_data.shape[0] / max(msk.shape[0], 1)
            zx = map_data.shape[1] / max(msk.shape[1], 1)
            msk = zoom(msk.astype(float), (zy, zx), order=1) > 0.5
        vals = map_data[msk]
        finite = vals[np.isfinite(vals)]
        if finite.size == 0:
            return None
        return float(np.mean(finite)), float(np.std(finite)), int(finite.size)
    except Exception:
        return None


def _agreement_stats(x: np.ndarray, y: np.ndarray) -> dict | None:
    """Regression / agreement metrics between paired samples x, y.

    Returns dict with n, slope, intercept (OLS y = slope·x + intercept),
    Pearson r, R² (= r²), Lin's concordance correlation coefficient (ρc) and
    the sample covariance — or None if fewer than two finite pairs.
    """
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    m = np.isfinite(x) & np.isfinite(y)
    x, y = x[m], y[m]
    n = int(x.size)
    if n < 2:
        return None
    xbar, ybar = float(x.mean()), float(y.mean())
    # population (1/n) central moments — Lin's CCC is defined with these
    sx2 = float(np.mean((x - xbar) ** 2))
    sy2 = float(np.mean((y - ybar) ** 2))
    sxy = float(np.mean((x - xbar) * (y - ybar)))
    sx, sy = np.sqrt(sx2), np.sqrt(sy2)
    pearson = sxy / (sx * sy) if sx > 0 and sy > 0 else float('nan')
    slope = sxy / sx2 if sx2 > 0 else float('nan')
    intercept = ybar - slope * xbar if np.isfinite(slope) else float('nan')
    denom = sx2 + sy2 + (xbar - ybar) ** 2
    lccc = (2.0 * sxy / denom) if denom > 0 else float('nan')
    cov_sample = float(np.cov(x, y, ddof=1)[0, 1]) if n >= 2 else float('nan')
    return {
        "n": n, "slope": slope, "intercept": intercept,
        "pearson": pearson, "r2": pearson ** 2 if np.isfinite(pearson) else float('nan'),
        "lccc": lccc, "covariance": cov_sample,
    }


def _roi_pixels(roi, map_data: np.ndarray) -> np.ndarray:
    """Finite pixel values of ``map_data`` inside ``roi.mask`` (mask auto-resized)."""
    try:
        msk = roi.mask
        if msk.shape != map_data.shape[:2]:
            from scipy.ndimage import zoom
            zy = map_data.shape[0] / max(msk.shape[0], 1)
            zx = map_data.shape[1] / max(msk.shape[1], 1)
            msk = zoom(msk.astype(float), (zy, zx), order=1) > 0.5
        vals = np.asarray(map_data)[msk]
        return vals[np.isfinite(vals)]
    except Exception:
        return np.asarray([], dtype=float)


def _parse_label(label: str) -> tuple[str, str]:
    """
    Split 'Method: pool' labels into (method, pool).
    For labels without ': ' there is no pool grouping — the whole label is the
    method (column header), so return (label, '') rather than a dummy 'Other'.
    """
    if ': ' in label:
        method, pool = label.split(': ', 1)
        return method.strip(), pool.strip()
    return label, ''


def _fmt_cell(mean: float, std: float) -> str:
    """Return compact 'mean ± std' string."""
    mag = abs(mean)
    if mag == 0 or (1e-3 <= mag < 1e4):
        return f"{mean:.4g} ± {std:.3g}"
    return f"{mean:.3e} ± {std:.2e}"


def _item(text: str,
          align: Qt.AlignmentFlag = Qt.AlignmentFlag.AlignCenter,
          bg: str | None = None,
          fg: str = "#e0e0e0",
          bold: bool = False,
          selectable: bool = True) -> QTableWidgetItem:
    it = QTableWidgetItem(text)
    it.setTextAlignment(align)
    if bg:
        it.setBackground(QBrush(QColor(bg)))
    it.setForeground(QBrush(QColor(fg)))
    if bold:
        f = QFont(); f.setBold(True); it.setFont(f)
    flags = Qt.ItemFlag.ItemIsEnabled
    if selectable:
        flags |= Qt.ItemFlag.ItemIsSelectable
    it.setFlags(flags)
    return it


# ─────────────────────────────────────────────────────────────────────────────
# Dialog
# ─────────────────────────────────────────────────────────────────────────────

class RoiTableDialog(QDialog):
    """
    Interactive pivot-style ROI statistics table.

    Columns are grouped by POOL (outer) × METHOD (inner).
    Each data cell shows  Mean ± Std.
    Checkboxes filter which pools / methods are visible (live rebuild).
    """

    def __init__(self, parent, rois: list,
                 map_items: list[tuple[str, np.ndarray]],
                 title: str = "ROI Statistics"):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setWindowFlags(
            self.windowFlags()
            | Qt.WindowType.WindowMaximizeButtonHint
            | Qt.WindowType.WindowMinimizeButtonHint
        )
        self.resize(1200, 620)

        self._rois      = rois
        self._map_items = map_items

        # ── Parse labels → (method, pool, arr) ───────────────────────────────
        self._entries: list[tuple[str, str, np.ndarray]] = []
        for label, arr in map_items:
            method, pool = _parse_label(label)
            self._entries.append((method, pool, arr))

        # Unique pools / methods preserving encounter order
        self._all_pools   = list(dict.fromkeys(p for _, p, _ in self._entries))
        self._all_methods = list(dict.fromkeys(m for m, _, _ in self._entries))

        # Quick lookup: (pool, method) → arr
        self._idx: dict[tuple[str, str], np.ndarray] = {
            (p, m): a for m, p, a in self._entries
        }

        # ── Outer layout ─────────────────────────────────────────────────────
        vl = QVBoxLayout(self)
        vl.setSpacing(4)

        # ── Filter bar ────────────────────────────────────────────────────────
        filter_frame = QFrame()
        filter_frame.setFrameShape(QFrame.Shape.StyledPanel)
        filter_frame.setStyleSheet(
            "QFrame{background:#1e1e2e;border-radius:5px;padding:2px;}"
            "QLabel{color:#ccc;font-size:11px;}"
            "QCheckBox{color:#ddd;font-size:11px;}"
            "QPushButton{font-size:10px;padding:1px 6px;}"
        )
        ff_vl = QVBoxLayout(filter_frame)
        ff_vl.setContentsMargins(8, 4, 8, 4)
        ff_vl.setSpacing(3)

        # One decluttered filter line (pool checkboxes · method checkboxes ·
        # Mean±SD) — no "Pools:"/"Methods:"/"Display:" labels.  A flow layout
        # keeps it on one line when there are few items and wraps only if many.
        flow = _FlowLayout(spacing=12)

        def _sep():
            s = QLabel("│"); s.setStyleSheet("color:#555;")
            return s

        self._pool_chks: dict[str, QCheckBox] = {}
        for pool in self._all_pools:
            if pool == '':                 # no-pool quantity → no filter checkbox
                continue
            chk = QCheckBox(pool); chk.setChecked(True)
            chk.toggled.connect(self._rebuild)
            self._pool_chks[pool] = chk
            flow.addWidget(chk)
        if len(self._all_pools) > 1:
            _ba_p = QPushButton("All");  _ba_p.setFixedWidth(36)
            _bn_p = QPushButton("None"); _bn_p.setFixedWidth(44)
            _ba_p.clicked.connect(lambda: [c.setChecked(True)  for c in self._pool_chks.values()])
            _bn_p.clicked.connect(lambda: [c.setChecked(False) for c in self._pool_chks.values()])
            flow.addWidget(_ba_p); flow.addWidget(_bn_p)

        flow.addWidget(_sep())

        self._meth_chks: dict[str, QCheckBox] = {}
        for meth in self._all_methods:
            chk = QCheckBox(meth); chk.setChecked(True)
            chk.toggled.connect(self._rebuild)
            self._meth_chks[meth] = chk
            flow.addWidget(chk)
        if len(self._all_methods) > 1:
            _ba_m = QPushButton("All");  _ba_m.setFixedWidth(36)
            _bn_m = QPushButton("None"); _bn_m.setFixedWidth(44)
            _ba_m.clicked.connect(lambda: [c.setChecked(True)  for c in self._meth_chks.values()])
            _bn_m.clicked.connect(lambda: [c.setChecked(False) for c in self._meth_chks.values()])
            flow.addWidget(_ba_m); flow.addWidget(_bn_m)

        flow.addWidget(_sep())

        self._chk_combined = QCheckBox("Mean ± SD")
        self._chk_combined.setChecked(True)
        self._chk_combined.setToolTip(
            "Checked: each cell shows 'Mean ± SD' in one column.\n"
            "Unchecked: Mean and SD are split into two separate columns.")
        self._chk_combined.toggled.connect(self._rebuild)
        flow.addWidget(self._chk_combined)

        ff_vl.addLayout(flow)

        vl.addWidget(filter_frame)

        # ── Info label ────────────────────────────────────────────────────────
        self._info_lbl = QLabel("")
        self._info_lbl.setStyleSheet("font-size:11px;color:#aaa;padding:2px;")
        vl.addWidget(self._info_lbl)

        # ── Table area ────────────────────────────────────────────────────────
        self._tbl_area = QWidget()
        self._tbl_area.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self._tbl_vl = QVBoxLayout(self._tbl_area)
        self._tbl_vl.setContentsMargins(0, 0, 0, 0)
        self._current_wgt: QWidget | None = None
        vl.addWidget(self._tbl_area, stretch=1)

        # ── Button row ────────────────────────────────────────────────────────
        btn_hl = QHBoxLayout()
        btn_clip = QPushButton("Copy to Clipboard (CSV)")
        btn_clip.clicked.connect(self._copy_csv)
        btn_hl.addWidget(btn_clip)

        btn_xlsx = QPushButton("Export to Excel (.xlsx)…")
        btn_xlsx.clicked.connect(self._export_xlsx)
        btn_hl.addWidget(btn_xlsx)

        btn_reg = QPushButton("X-Y Plot…")
        btn_reg.clicked.connect(self._open_regression)
        btn_hl.addWidget(btn_reg)

        btn_box = QPushButton("Box plots…")
        btn_box.clicked.connect(self._open_boxplots)
        btn_hl.addWidget(btn_box)

        btn_hl.addStretch()
        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.accept)
        btn_hl.addWidget(btn_close)
        vl.addLayout(btn_hl)

        # ── Initial build ─────────────────────────────────────────────────────
        self._rebuild()

    # ─────────────────────────────────────────────────────────────────────────
    # Helpers
    # ─────────────────────────────────────────────────────────────────────────

    def _active_pools(self)   -> list[str]:
        return [p for p in self._all_pools
                if p == '' or self._pool_chks[p].isChecked()]

    def _active_methods(self) -> list[str]:
        return [m for m in self._all_methods if self._meth_chks[m].isChecked()]

    # ─────────────────────────────────────────────────────────────────────────
    # Table builder
    # ─────────────────────────────────────────────────────────────────────────

    def _rebuild(self):
        """Destroy old table and build a fresh one from current filter state."""
        # Remove whatever widget currently occupies the table area
        if self._current_wgt is not None:
            self._tbl_vl.removeWidget(self._current_wgt)
            self._current_wgt.deleteLater()
            self._current_wgt = None

        act_pools   = self._active_pools()
        act_methods = self._active_methods()

        # (pool, method) pairs that actually have data — ordered pool-first
        col_defs: list[tuple[str, str]] = [
            (pool, meth)
            for pool in act_pools
            for meth in act_methods
            if (pool, meth) in self._idx
        ]

        if not col_defs:
            lbl = QLabel("No data for selected pools / methods.")
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl.setStyleSheet("color:#888;font-size:13px;padding:40px;")
            self._tbl_vl.addWidget(lbl)
            self._current_wgt = lbl
            self._info_lbl.setText("Select at least one pool and one method above.")
            return

        # ── Combined 'Mean ± SD' (one col) vs split Mean / SD (two cols) ──────
        combined = self._chk_combined.isChecked()
        # Physical columns: (pool, method, kind)  kind ∈ {"ms", "mean", "sd"}
        phys_cols: list[tuple[str, str, str]] = []
        for pool, meth in col_defs:
            if combined:
                phys_cols.append((pool, meth, "ms"))
            else:
                phys_cols.append((pool, meth, "mean"))
                phys_cols.append((pool, meth, "sd"))
        self._phys_cols = phys_cols

        n_rois    = len(self._rois)
        n_data    = len(phys_cols)
        # Total columns: ROI name | (pool×method data) | n pixels
        n_cols    = 1 + n_data + 1
        # Total rows: pool-header row + ROI data rows
        n_rows    = 1 + n_rois

        self._info_lbl.setText(
            f"  {n_rois} ROI(s)  ×  {len(col_defs)} column(s)   —   "
            + ("values: Mean ± Std" if combined
               else "values: separate Mean and SD columns")
        )

        tbl = QTableWidget(n_rows, n_cols)
        tbl.setAlternatingRowColors(False)
        tbl.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        tbl.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)
        tbl.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        # Track the ORDER in which ROI rows are selected, so copy preserves it
        self._sel_order: list = []
        tbl.itemSelectionChanged.connect(self._on_table_selection_changed)
        # Ctrl+C / Cmd+C → copy (selection-aware)
        from PyQt6.QtGui import QShortcut, QKeySequence
        _sc = QShortcut(QKeySequence.StandardKey.Copy, tbl)
        _sc.activated.connect(self._copy_csv)
        tbl.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents)
        tbl.horizontalHeader().setStretchLastSection(False)
        tbl.verticalHeader().setVisible(False)
        tbl.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)

        # ── Real column headers (method names, row -1) ────────────────────────
        # Column 0: "ROI"  |  cols 1..N_data: method  |  col N+1: "n (pixels)"
        h_labels = ["ROI"]
        for _p, meth, kind in phys_cols:
            if kind == "ms":
                h_labels.append(meth)
            elif kind == "mean":
                h_labels.append(f"{meth}\nMean")
            else:
                h_labels.append(f"{meth}\nSD")
        h_labels.append("n\n(pixels)")
        tbl.setHorizontalHeaderLabels(h_labels)

        # ── Pool colour map ───────────────────────────────────────────────────
        pool_ci = {p: i % len(_POOL_BG) for i, p in enumerate(act_pools)}

        # ── Theme-adaptive colours (white text on dark OS, black on light OS) ─
        _is_dark   = _is_dark_theme()
        _blank_bg  = "#2b2b3b" if _is_dark else "#f0f0f0"

        # ── Row 0: pool-group header (spans over all method cols per pool) ────
        # Col 0: blank header
        tbl.setItem(0, 0, _item("", bg=_blank_bg, selectable=False))
        # Col N+1: blank header
        tbl.setItem(0, n_cols - 1, _item("", bg=_blank_bg, selectable=False))

        # Build per-pool start/end column indices (1-based, over physical cols)
        pool_spans: dict[str, tuple[int, int]] = {}
        cur = 1
        for pool in act_pools:
            cnt = sum(1 for p, _m, _k in phys_cols if p == pool)
            if cnt == 0:
                continue
            pool_spans[pool] = (cur, cur + cnt - 1)
            cur += cnt

        pool_hdr_font = QFont()
        pool_hdr_font.setBold(True)
        pool_hdr_font.setPointSize(13)

        for pool, (c0, c1) in pool_spans.items():
            if pool == '':                          # no-pool quantity → blank bar
                bg, fg, _txt = _blank_bg, "#888888", ""
            else:
                ci      = pool_ci[pool]
                bg, fg  = _pool_header_colors(_POOL_BG[ci % len(_POOL_BG)], _is_dark)
                _txt    = _display_pool_name(pool)
            it      = QTableWidgetItem(_txt)
            it.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            it.setFont(pool_hdr_font)
            it.setBackground(QBrush(QColor(bg)))
            it.setForeground(QBrush(QColor(fg)))
            flags = Qt.ItemFlag.ItemIsEnabled   # NOT selectable
            it.setFlags(flags)
            tbl.setItem(0, c0, it)
            if c1 > c0:
                tbl.setSpan(0, c0, 1, c1 - c0 + 1)

        tbl.setRowHeight(0, 30)

        # ── Data rows (rows 1 … n_rois) ───────────────────────────────────────
        bold_f = QFont(); bold_f.setBold(True)

        self._raw_data: list[list]  = []
        self._col_defs              = col_defs
        # Export headers aligned to PHYSICAL columns (combined or split)
        _exp = ["ROI"]
        for pool, meth, kind in phys_cols:
            _pref = f"{pool}: " if pool else ""
            if kind == "ms":
                _exp.append(f"{_pref}{meth}")
            elif kind == "mean":
                _exp.append(f"{_pref}{meth} (mean)")
            else:
                _exp.append(f"{_pref}{meth} (SD)")
        _exp.append("n_pixels")
        self._export_hdrs = _exp

        _even_row = "#18181e"
        _odd_row  = "#141418"

        def _fmt_num(v: float) -> str:
            return (f"{v:.4g}" if (v == 0 or 1e-3 <= abs(v) < 1e4) else f"{v:.3e}")

        for ri, roi in enumerate(self._rois):
            tr      = ri + 1                     # table row (shifted by pool-header row)
            row_bg  = _even_row if ri % 2 == 0 else _odd_row
            row_exp = [roi.name]

            # ROI name cell (left-aligned, bold)
            roi_it = QTableWidgetItem(roi.name)
            roi_it.setFont(bold_f)
            roi_it.setTextAlignment(
                Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft)
            roi_it.setBackground(QBrush(QColor("#22222e")))
            roi_it.setForeground(QBrush(QColor("#e8e8ff")))
            roi_it.setFlags(
                Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
            tbl.setItem(tr, 0, roi_it)

            n_pix: int | None = None
            _stat_cache: dict[tuple[str, str], tuple | None] = {}

            for ci, (pool, meth, kind) in enumerate(phys_cols):
                col = ci + 1
                key = (pool, meth)
                if key not in _stat_cache:
                    arr = self._idx.get(key)
                    _stat_cache[key] = _roi_stats(roi, arr) if arr is not None else None
                result = _stat_cache[key]

                if result is not None:
                    mean_v, std_v, n_v = result
                    if n_pix is None:
                        n_pix = n_v
                    if kind == "ms":
                        text = _fmt_cell(mean_v, std_v)
                        row_exp.append(f"{mean_v:.6g}±{std_v:.6g}")
                    elif kind == "mean":
                        text = _fmt_num(mean_v)
                        row_exp.append(f"{mean_v:.6g}")
                    else:  # "sd"
                        text = _fmt_num(std_v)
                        row_exp.append(f"{std_v:.6g}")
                else:
                    text = "—"
                    row_exp.append(None)

                tbl.setItem(tr, col, _item(text, bg=row_bg))

            # n pixels (last column)
            n_col = n_cols - 1
            if n_pix is not None:
                tbl.setItem(tr, n_col, _item(str(n_pix), bg="#1e1e28"))
                row_exp.append(n_pix)
            else:
                tbl.setItem(tr, n_col, _item("—", bg="#1e1e28"))
                row_exp.append(None)

            self._raw_data.append(row_exp)

        self._current_wgt = tbl
        self._tbl_vl.addWidget(tbl)

    # ─────────────────────────────────────────────────────────────────────────
    # Export
    # ─────────────────────────────────────────────────────────────────────────

    def _on_table_selection_changed(self):
        """Maintain ``self._sel_order`` = data-row indices in selection order."""
        wgt = getattr(self, '_current_wgt', None)
        if wgt is None:
            return
        cur = {it.row() - 1 for it in wgt.selectedItems()
               if it.row() >= 1 and (it.row() - 1) < len(self._raw_data or [])}
        # Drop deselected rows, keep existing order; append newly-selected rows
        self._sel_order = [r for r in self._sel_order if r in cur]
        for r in sorted(cur):                       # added-this-event → table order
            if r not in self._sel_order:
                self._sel_order.append(r)

    def _copy_csv(self):
        if not getattr(self, '_raw_data', None):
            return
        from PyQt6.QtWidgets import QApplication

        # Selection-aware copy: only the selected ROWS and COLUMNS are copied.
        # (Table row 0 = pool-group header → data row r maps to _raw_data[r-1];
        #  table column c maps directly to _export_hdrs[c] / _raw_data[r][c].)
        wgt   = getattr(self, '_current_wgt', None)
        items = wgt.selectedItems() if wgt is not None else []
        n_cols = len(self._export_hdrs)
        sel_rows = {it.row() - 1 for it in items
                    if it.row() >= 1 and 0 <= it.row() - 1 < len(self._raw_data)}
        sel_cols = {it.column() for it in items
                    if it.row() >= 1 and 0 <= it.column() < n_cols}

        if sel_rows and sel_cols:
            # Rows in the order they were selected; then any others in table order
            rows_idx = [r for r in getattr(self, '_sel_order', []) if r in sel_rows]
            for r in sorted(sel_rows):
                if r not in rows_idx:
                    rows_idx.append(r)
            cols_idx = sorted(sel_cols)
            note = f"  ✓ Copied {len(rows_idx)} ROI(s) × {len(cols_idx)} column(s)!"
        else:
            # Nothing selected → copy the whole table
            rows_idx = list(range(len(self._raw_data)))
            cols_idx = list(range(n_cols))
            note = "  ✓ Copied all ROIs!"

        lines = ["\t".join(self._export_hdrs[c] for c in cols_idx)]
        for r in rows_idx:
            row = self._raw_data[r]
            lines.append("\t".join(
                "" if row[c] is None else str(row[c]) for c in cols_idx))
        QApplication.clipboard().setText("\n".join(lines))
        old = self.windowTitle()
        self.setWindowTitle(old + note)
        from PyQt6.QtCore import QTimer
        QTimer.singleShot(1800, lambda: self.setWindowTitle(old))

    def _export_xlsx(self):
        if not getattr(self, '_raw_data', None):
            return
        from PyQt6.QtWidgets import QFileDialog
        path, _ = QFileDialog.getSaveFileName(
            self, "Export ROI table", "roi_stats.xlsx",
            "Excel files (*.xlsx);;All files (*)"
        )
        if not path:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "ROI Stats"

            phys_cols = getattr(self, "_phys_cols",
                                [(p, m, "ms") for p, m in self._col_defs])

            # ── Row 1: Pool group headers (merged over physical columns) ──────
            ws.append(["ROI"] + [None] * len(phys_cols) + ["n (pixels)"])
            cur_xl = 2   # Excel column index starts at 1; col A = ROI
            for pool in self._active_pools():
                cnt = sum(1 for p, _m, _k in phys_cols if p == pool)
                if cnt == 0:
                    continue
                c_start = cur_xl
                c_end   = cur_xl + cnt - 1
                cell = ws.cell(1, c_start, pool)
                cell.font      = Font(bold=True, color="FFFFFF", size=11)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill      = PatternFill("solid", fgColor="1a3a5c")
                if c_end > c_start:
                    ws.merge_cells(
                        start_row=1, start_column=c_start,
                        end_row=1,   end_column=c_end)
                cur_xl = c_end + 1

            # ── Row 2: Method sub-headers (physical columns) ──────────────────
            def _sub(meth, kind):
                return meth if kind == "ms" else (
                    f"{meth} (mean)" if kind == "mean" else f"{meth} (SD)")
            sub_hdrs = (["ROI"]
                        + [_sub(m, k) for _p, m, k in phys_cols]
                        + ["n (pixels)"])
            ws.append(sub_hdrs)
            for cell in ws[2]:
                cell.font      = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center")
                cell.fill      = PatternFill("solid", fgColor="2a2a4a")

            # ── Data rows ─────────────────────────────────────────────────────
            for row in self._raw_data:
                ws.append([("" if v is None else v) for v in row])

            # Auto column widths
            for col in ws.columns:
                max_w = max((len(str(cell.value or "")) for cell in col), default=6)
                ws.column_dimensions[
                    get_column_letter(col[0].column)
                ].width = min(max_w + 4, 30)

            ws.freeze_panes = "B3"   # freeze ROI column + first two header rows

            wb.save(path)
            QMessageBox.information(self, "Saved", f"Exported to:\n{path}")

        except ImportError:
            QMessageBox.warning(
                self, "Missing dependency",
                "openpyxl is not installed.\n"
                "Install with:  pip install openpyxl\n\n"
                "You can still use 'Copy to Clipboard' and paste into Excel."
            )
        except Exception as exc:
            QMessageBox.critical(self, "Export error", str(exc))

    # ─────────────────────────────────────────────────────────────────────────
    # Regression / agreement between two columns (method comparison)
    # ─────────────────────────────────────────────────────────────────────────

    def _open_regression(self):
        # Always open — the X-Y dialog also offers the ROI index and pixel count
        # as axes, so a plot is possible even with a single data column.
        cols = list(self._idx.keys())          # [(pool, method), …]
        dlg = _RegressionDialog(self, self._rois, self._idx, cols)
        dlg.show()

    def _open_boxplots(self):
        cols = [(p, m) for p in self._active_pools() for m in self._active_methods()
                if (p, m) in self._idx]
        if not cols:
            cols = list(self._idx.keys())
        dlg = _BoxPlotDialog(self, self._rois, self._idx, cols)
        dlg.show()


# ─────────────────────────────────────────────────────────────────────────────
# Regression / agreement dialog
# ─────────────────────────────────────────────────────────────────────────────

class _PlotStyleBar(QWidget):
    """Compact font-size + colour toolbar shared by the ROI-statistics plots:
    Title / Main / Axes / Ticks / Legend sizes, a dark 'Bg' toggle, and a single
    custom plot colour.  Emits ``changed`` (debounced for the spinboxes) whenever
    any control is touched, so the owning dialog can re-render."""

    changed = pyqtSignal()

    def __init__(self, parent=None, defaults=None):
        super().__init__(parent)
        d = {"title": 11, "main": 13, "axes": 10, "ticks": 9, "legend": 9}
        if defaults:
            d.update(defaults)
        row = QHBoxLayout(self)
        row.setContentsMargins(2, 2, 2, 2)
        row.setSpacing(6)

        def _sp(label, key):
            row.addWidget(QLabel(label))
            s = QSpinBox(); s.setRange(4, 40); s.setValue(d[key]); s.setFixedWidth(48)
            row.addWidget(s)
            return s

        self.sp_title  = _sp("Title:",  "title")
        self.sp_main   = _sp("Main:",   "main")
        self.sp_axes   = _sp("Axes:",   "axes")
        self.sp_ticks  = _sp("Ticks:",  "ticks")
        self.sp_legend = _sp("Legend:", "legend")

        self.chk_bg = QCheckBox("Bg")
        self.chk_bg.setToolTip("Black background for the figure (for slides).")
        row.addWidget(self.chk_bg)
        self.chk_color = QCheckBox("Custom color")
        row.addWidget(self.chk_color)
        self.btn_color = QPushButton(); self.btn_color.setFixedWidth(30)
        self._color = "#1f77b4"; self._swatch()
        row.addWidget(self.btn_color)
        row.addStretch()

        self._timer = QTimer(self)
        self._timer.setSingleShot(True); self._timer.setInterval(160)
        self._timer.timeout.connect(self.changed.emit)
        for s in (self.sp_title, self.sp_main, self.sp_axes,
                  self.sp_ticks, self.sp_legend):
            s.valueChanged.connect(lambda _: self._timer.start())
        self.chk_bg.toggled.connect(lambda _: self.changed.emit())
        self.chk_color.toggled.connect(lambda _: self.changed.emit())
        self.btn_color.clicked.connect(self._pick)

    def _swatch(self):
        self.btn_color.setStyleSheet(
            f"background:{self._color};border:1px solid #888;border-radius:3px;")

    def _pick(self):
        c = QColorDialog.getColor(QColor(self._color), self, "Plot colour")
        if c.isValid():
            self._color = c.name(); self._swatch()
            if self.chk_color.isChecked():
                self.changed.emit()

    def sizes(self) -> dict:
        return dict(title=self.sp_title.value(), main=self.sp_main.value(),
                    axes=self.sp_axes.value(), ticks=self.sp_ticks.value(),
                    legend=self.sp_legend.value())

    def bg(self) -> bool:
        return self.chk_bg.isChecked()

    def color(self):
        """Chosen colour hex when 'Custom color' is ticked, else None."""
        return self._color if self.chk_color.isChecked() else None


def _sig_stars(p) -> str:
    """Significance marker: ** for p<0.005, * for p<0.05, else 'ns'."""
    if p is None or not np.isfinite(p):
        return ""
    if p < 0.005:
        return "**"
    if p < 0.05:
        return "*"
    return "ns"


def _group_compare(groups: list):
    """Compare 2+ groups (each a 1-D array of values).  Picks the test by group
    count and normality (Shapiro–Wilk on each group; parametric only if every
    group looks normal):
        2 groups  → Welch t-test (normal)  |  Mann-Whitney U (non-parametric)
        3+ groups → one-way ANOVA (normal) |  Kruskal-Wallis (non-parametric)
    Returns dict(test, stat, p, n_groups, parametric) or None if <2 usable groups."""
    from scipy import stats as _st
    gs = []
    for g in groups:
        a = np.asarray(g, float)
        a = a[np.isfinite(a)]
        if a.size >= 2:
            gs.append(a)
    if len(gs) < 2:
        return None
    parametric = True
    for a in gs:
        if a.size >= 3 and float(np.ptp(a)) > 0:
            try:
                if _st.shapiro(a)[1] < 0.05:
                    parametric = False; break
            except Exception:
                parametric = False; break
        else:
            parametric = False; break
    try:
        if len(gs) == 2:
            if parametric:
                stat, p = _st.ttest_ind(gs[0], gs[1], equal_var=False)
                name = "Welch t-test"
            else:
                stat, p = _st.mannwhitneyu(gs[0], gs[1], alternative="two-sided")
                name = "Mann-Whitney U"
        else:
            if parametric:
                stat, p = _st.f_oneway(*gs)
                name = "One-way ANOVA"
            else:
                stat, p = _st.kruskal(*gs)
                name = "Kruskal-Wallis"
    except Exception:
        return None
    return dict(test=name, stat=float(stat), p=float(p),
                n_groups=len(gs), parametric=parametric)


class _RegressionDialog(QDialog):
    """Correlate two ROI-statistic columns across ROIs and report regression /
    agreement metrics (slope, intercept, Pearson r, R², Lin's CCC, covariance).
    Each ROI contributes one paired point: its mean in the X map vs the Y map."""

    def __init__(self, parent, rois, idx, cols):
        super().__init__(parent)
        from PyQt6.QtWidgets import QComboBox
        self.setWindowTitle("X-Y Plot")
        self.resize(540, 480)
        self._rois = rois
        self._idx  = idx
        self._cols = cols

        def _lab(pm): return pm[1] if not pm[0] else f"{pm[0]}: {pm[1]}"

        # Selectable axes = every data column + the ROI index + the pixel count,
        # so an X-Y plot is possible even with a single data column.
        self._axes: list[tuple[str, object]] = []
        for _c in cols:
            self._axes.append((_lab(_c),
                               (lambda roi, i, c=_c: self._col_mean(roi, c))))
        self._axes.append(("ROI index", (lambda roi, i: float(i + 1))))
        self._axes.append(("n (pixels)", (lambda roi, i: self._col_n(roi))))
        # User-entered values (one per ROI) — assignable to either X or Y.
        self._custom_vals = [np.nan] * len(self._rois)
        self._axes.append(("Custom values", (lambda roi, i: self._custom_vals[i])))
        _labels = [a[0] for a in self._axes]

        v = QVBoxLayout(self)
        row = QHBoxLayout()
        row.addWidget(QLabel("X:"))
        self._cx = QComboBox(); self._cx.addItems(_labels)
        row.addWidget(self._cx, 1)
        row.addWidget(QLabel("Y:"))
        self._cy = QComboBox(); self._cy.addItems(_labels)
        self._cy.setCurrentIndex(1 if len(_labels) > 1 else 0)
        row.addWidget(self._cy, 1)
        v.addLayout(row)
        self._cx.currentIndexChanged.connect(self._compute)
        self._cy.currentIndexChanged.connect(self._compute)

        crow = QHBoxLayout()
        btn_custom = QPushButton("Set custom values…")
        btn_custom.setToolTip(
            "Enter one value per ROI (e.g. nominal concentration, pH, or a "
            "reference measurement), then choose 'Custom values' for X or Y.")
        btn_custom.clicked.connect(self._edit_custom)
        crow.addWidget(btn_custom); crow.addStretch()
        v.addLayout(crow)

        # Font-size / colour toolbar (re-renders the plot on change).
        self._style = _PlotStyleBar(self)
        self._style.changed.connect(self._compute)
        v.addWidget(self._style)

        # ── Scatter plot (X vs Y) with the fitted regression line ─────────────
        from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as _FC
        from matplotlib.figure import Figure as _Fig
        self._fig = _Fig(figsize=(4.6, 3.3), facecolor="white")
        self._ax = self._fig.add_subplot(111)
        self._canvas = _FC(self._fig)
        v.addWidget(self._canvas, 3)

        self._note = QLabel("")
        self._note.setWordWrap(True)
        self._note.setStyleSheet("color:#aaa;font-size:11px;padding:2px;")
        v.addWidget(self._note)

        # Stats numbers shown BELOW the plot.
        self._res = QTableWidget(0, 2)
        self._res.setHorizontalHeaderLabels(["Statistic", "Value"])
        self._res.horizontalHeader().setStretchLastSection(True)
        self._res.verticalHeader().setVisible(False)
        self._res.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._res.setMaximumHeight(230)
        v.addWidget(self._res, 1)

        brow = QHBoxLayout()
        btn_copy = QPushButton("Copy")
        btn_copy.clicked.connect(self._copy)
        brow.addWidget(btn_copy)
        btn_save = QPushButton("Save X-Y Plot…")
        btn_save.setToolTip("Save the plot as PNG / JPEG / PDF / SVG / TIFF, "
                            "or the paired X-Y data + statistics as Excel (.xlsx).")
        btn_save.clicked.connect(self._save)
        brow.addWidget(btn_save)
        brow.addStretch()
        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.accept)
        brow.addWidget(btn_close)
        v.addLayout(brow)

        self._rows: list[tuple[str, str]] = []
        self._compute()

    def _col_mean(self, roi, key):
        """ROI mean of a data column (or None if unavailable)."""
        arr = self._idx.get(key)
        if arr is None:
            return None
        r = _roi_stats(roi, arr)
        return r[0] if r is not None else None

    def _col_n(self, roi):
        """Pixel count for this ROI (from the first available data column)."""
        for c in self._cols:
            arr = self._idx.get(c)
            if arr is None:
                continue
            r = _roi_stats(roi, arr)
            if r is not None:
                return float(r[2])
        return None

    def _edit_custom(self):
        """Enter one value per ROI for the 'Custom values' axis."""
        d = QDialog(self)
        d.setWindowTitle("Custom values (one per ROI)")
        d.resize(320, 440)
        lay = QVBoxLayout(d)
        lay.addWidget(QLabel("Enter a value for each ROI (blank = skip):"))
        tbl = QTableWidget(len(self._rois), 2)
        tbl.setHorizontalHeaderLabels(["ROI", "Value"])
        tbl.verticalHeader().setVisible(False)
        tbl.horizontalHeader().setStretchLastSection(True)
        for i, roi in enumerate(self._rois):
            it = QTableWidgetItem(roi.name)
            it.setFlags(Qt.ItemFlag.ItemIsEnabled)
            tbl.setItem(i, 0, it)
            v0 = self._custom_vals[i]
            txt = "" if (v0 is None or not np.isfinite(v0)) else f"{v0:g}"
            tbl.setItem(i, 1, QTableWidgetItem(txt))
        tbl.resizeColumnsToContents()
        lay.addWidget(tbl, 1)
        brow = QHBoxLayout()
        b_ok = QPushButton("OK"); b_cancel = QPushButton("Cancel")
        brow.addStretch(); brow.addWidget(b_cancel); brow.addWidget(b_ok)
        lay.addLayout(brow)
        b_cancel.clicked.connect(d.reject); b_ok.clicked.connect(d.accept)
        if d.exec():
            for i in range(len(self._rois)):
                it = tbl.item(i, 1)
                s = it.text().strip() if it else ""
                try:
                    self._custom_vals[i] = float(s) if s else np.nan
                except ValueError:
                    self._custom_vals[i] = np.nan
            # If either axis is on 'Custom values', reflect the new data now.
            ci = self._axes[self._cx.currentIndex()][0]
            cj = self._axes[self._cy.currentIndex()][0]
            if "Custom values" in (ci, cj):
                self._compute()

    def _pair(self) -> tuple[np.ndarray, np.ndarray]:
        gx = self._axes[self._cx.currentIndex()][1]
        gy = self._axes[self._cy.currentIndex()][1]
        xs, ys = [], []
        for i, roi in enumerate(self._rois):
            rx = gx(roi, i); ry = gy(roi, i)
            if (rx is not None and ry is not None
                    and np.isfinite(rx) and np.isfinite(ry)):
                xs.append(float(rx)); ys.append(float(ry))
        return np.asarray(xs, float), np.asarray(ys, float)

    def _draw_plot(self, x, y, st):
        """Black-and-white scatter of X vs Y with the fitted regression line."""
        xl = self._axes[self._cx.currentIndex()][0]
        yl = self._axes[self._cy.currentIndex()][0]
        fs = self._style.sizes()
        _c = self._style.color() or "black"
        ax = self._ax
        ax.clear()
        ax.set_facecolor("white")
        if x.size:
            ax.scatter(x, y, facecolors="none", edgecolors=_c,
                       s=42, linewidths=1.3, zorder=3)
        if st is not None and np.isfinite(st["slope"]) and x.size >= 2:
            xx = np.array([float(np.min(x)), float(np.max(x))])
            _b = st["intercept"]
            _sgn = "+" if _b >= 0 else "−"
            # Legend: fit equation (slope + y-intercept) and R² only (no LCCC).
            ax.plot(xx, st["slope"] * xx + st["intercept"], "-",
                    color=_c, lw=1.6, zorder=2,
                    label=(f"y = {st['slope']:.3g}·x {_sgn} {abs(_b):.3g}\n"
                           f"R² = {st['r2']:.3f}"))
            ax.legend(fontsize=fs["legend"], framealpha=0.9,
                      edgecolor="0.6", loc="best")
        ax.set_xlabel(xl, fontsize=fs["axes"])
        ax.set_ylabel(yl, fontsize=fs["axes"])
        ax.set_title(f"{yl}  vs  {xl}", fontsize=fs["title"])
        ax.tick_params(labelsize=fs["ticks"])
        ax.grid(True, color="0.85", lw=0.6)
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        self._fig.tight_layout()
        try:
            from my_gui.fig_theme import apply_fig_dark_theme
            apply_fig_dark_theme(self._fig, self._style.bg())
        except Exception:
            pass
        self._canvas.draw()

    def _compute(self):
        x, y = self._pair()
        st = _agreement_stats(x, y)
        self._res.setRowCount(0)
        if st is None:
            self._note.setText(
                "Not enough paired ROIs — need at least two ROIs with a finite "
                "value in both selected axes.")
            self._note.setVisible(True)
            self._rows = []
            self._draw_plot(x, y, None)
            return
        self._note.setText("")
        self._note.setVisible(False)
        self._rows = [
            ("N (paired ROIs)",     f"{st['n']}"),
            ("Slope",               f"{st['slope']:.6g}"),
            ("Intercept",           f"{st['intercept']:.6g}"),
            ("Pearson r",           f"{st['pearson']:.6g}"),
            ("R²  (correlation)",   f"{st['r2']:.6g}"),
            ("Lin's CCC (LCCC)",    f"{st['lccc']:.6g}"),
            ("Covariance (sample)", f"{st['covariance']:.6g}"),
        ]
        self._res.setRowCount(len(self._rows))
        bold = QFont(); bold.setBold(True)
        for i, (k, val) in enumerate(self._rows):
            it_k = QTableWidgetItem(k); it_k.setFont(bold)
            self._res.setItem(i, 0, it_k)
            self._res.setItem(i, 1, QTableWidgetItem(val))
        self._draw_plot(x, y, st)

    def _copy(self):
        from PyQt6.QtWidgets import QApplication
        if not self._rows:
            return
        QApplication.clipboard().setText(
            "\n".join(f"{k}\t{v}" for k, v in self._rows))

    def _save(self):
        """Save the plot as an image, or the paired data + stats as .xlsx."""
        from my_gui.fig_export import save_figure, FIG_EXPORT_FILTER
        flt = FIG_EXPORT_FILTER + ";;Excel spreadsheet (*.xlsx)"
        p, _ = QFileDialog.getSaveFileName(self, "Save X-Y Plot", "xy_plot", flt)
        if not p:
            return
        if p.lower().endswith(".xlsx"):
            self._save_xlsx(p)
        else:
            save_figure(self._fig, p, dpi=300, facecolor="white")

    def _save_xlsx(self, path):
        try:
            from openpyxl import Workbook
        except Exception:
            QMessageBox.warning(self, "openpyxl missing",
                                "Excel export needs openpyxl  (pip install openpyxl).")
            return
        xl = self._axes[self._cx.currentIndex()][0]
        yl = self._axes[self._cy.currentIndex()][0]
        gx = self._axes[self._cx.currentIndex()][1]
        gy = self._axes[self._cy.currentIndex()][1]
        wb = Workbook()
        ws = wb.active; ws.title = "XY data"
        ws.append(["ROI", xl, yl])
        for i, roi in enumerate(self._rois):
            rx = gx(roi, i); ry = gy(roi, i)
            ws.append([roi.name,
                       float(rx) if rx is not None and np.isfinite(rx) else None,
                       float(ry) if ry is not None and np.isfinite(ry) else None])
        if self._rows:
            ws2 = wb.create_sheet("Statistics")
            ws2.append(["Statistic", "Value"])
            for k, val in self._rows:
                ws2.append([k, val])
        try:
            wb.save(path)
        except Exception as exc:
            QMessageBox.warning(self, "Save failed", str(exc))


# ─────────────────────────────────────────────────────────────────────────────
# Box-and-whisker plot dialog (black & white)
# ─────────────────────────────────────────────────────────────────────────────

class _BoxPlotDialog(QDialog):
    """Black-and-white box-and-whisker plots of each ROI's value distribution
    within a map — one panel per selected column."""

    def __init__(self, parent, rois, idx, cols):
        super().__init__(parent)
        self.setWindowTitle("ROI Box Plots")
        # Size the window to the number of map columns so every panel has room
        # from the first render (the figure tracks the widget size).
        _n = max(len(cols), 1)
        _nc = 1 if _n == 1 else 2
        _nr = (_n + _nc - 1) // _nc
        self.resize(max(640, 470 * _nc), max(460, 360 * _nr + 110))
        from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as _FC
        from matplotlib.figure import Figure as _Fig

        self._rois = list(rois)
        self._idx  = idx
        self._cols = list(cols)
        # Which ROIs to show — all selected by default (the "Select ROIs…" button
        # lets the user restrict the box plots to specific ROIs).
        self._selected = list(range(len(self._rois)))
        self._groups: dict = {}     # roi index -> group label (for statistics)
        self._stats: dict = {}      # (pool, meth) -> group-compare result dict

        v = QVBoxLayout(self)
        self._style = _PlotStyleBar(self)
        self._style.changed.connect(self._redraw)
        v.addWidget(self._style)
        self._fig = _Fig(facecolor="white")
        self._canvas = _FC(self._fig)
        v.addWidget(self._canvas, 1)

        brow = QHBoxLayout()
        btn_pick = QPushButton("Select ROIs…")
        btn_pick.setToolTip("Choose which ROIs appear in the box plots "
                            "(default: all).")
        btn_pick.clicked.connect(self._pick_rois)
        brow.addWidget(btn_pick)
        btn_save = QPushButton("Save figure…")
        btn_save.clicked.connect(self._save)
        brow.addWidget(btn_save)
        btn_stats = QPushButton("Statistics…")
        btn_stats.setToolTip(
            "Assign ROIs to groups and compare them — t-test / Mann-Whitney for "
            "2 groups, ANOVA / Kruskal-Wallis for 3+.  Significant results are "
            "starred on the plots (* p<0.05, ** p<0.005).")
        btn_stats.clicked.connect(self._run_stats)
        brow.addWidget(btn_stats)
        brow.addStretch()
        btn_close = QPushButton("Close"); btn_close.clicked.connect(self.accept)
        brow.addWidget(btn_close)
        v.addLayout(brow)

        self._redraw()

    def _redraw(self):
        sel = self._selected or list(range(len(self._rois)))
        rois = [self._rois[i] for i in sel]
        cols = self._cols
        fs = self._style.sizes()
        _boxc = self._style.color() or "white"
        self._fig.clf()
        n = max(len(cols), 1)
        ncol = 1 if n == 1 else 2
        nrow = (n + ncol - 1) // ncol
        # NOTE: do NOT set a fixed figure size here — the FigureCanvas already
        # tracks the widget size, and forcing set_size_inches() makes the render
        # mismatch the widget until a manual window resize fixes it.
        names = [r.name for r in rois]

        for ci, (pool, meth) in enumerate(cols):
            arr = self._idx.get((pool, meth))
            ax = self._fig.add_subplot(nrow, ncol, ci + 1)
            ax.set_facecolor("white")
            data = []
            for r in rois:
                px = _roi_pixels(r, arr) if arr is not None else np.asarray([], float)
                data.append(px if px.size else np.asarray([np.nan]))
            bp = ax.boxplot(data, patch_artist=True, showfliers=False, widths=0.6)
            for b in bp["boxes"]:
                b.set(facecolor=_boxc, edgecolor="black", linewidth=1.2)
            for w in bp["whiskers"] + bp["caps"]:
                w.set(color="black", linewidth=1.1)
            for m in bp["medians"]:
                m.set(color="black", linewidth=1.8)
            ax.set_xticks(range(1, len(names) + 1))
            ax.set_xticklabels(names, rotation=45, ha="right",
                               fontsize=fs["ticks"], color="black")
            ax.tick_params(axis="y", labelsize=fs["ticks"], colors="black")
            # Title = the map name only (no p-value text).
            ax.set_title(meth if not pool else f"{pool}: {meth}",
                         fontsize=fs["title"], color="black")
            ax.grid(True, axis="y", color="0.85", lw=0.6)
            for s in ("top", "right"):
                ax.spines[s].set_visible(False)
            for s in ("left", "bottom"):
                ax.spines[s].set_color("black")

            # Significance star ABOVE each grouped box (drawn only when the
            # group comparison for this column is significant).
            _st = self._stats.get((pool, meth))
            _stars = _sig_stars(_st["p"]) if _st else ""
            if _stars in ("*", "**"):
                ax.margins(y=0.14)          # headroom so the stars aren't clipped
                caps = bp["caps"]
                for bi in range(len(rois)):
                    if sel[bi] not in self._groups:
                        continue
                    try:
                        y_up = float(caps[2 * bi + 1].get_ydata()[0])
                    except Exception:
                        y_up = float(np.nanmax(data[bi]))
                    ax.text(bi + 1, y_up, _stars, ha="center", va="bottom",
                            fontsize=max(fs["title"] + 2, 13),
                            fontweight="bold", color="black")

        self._fig.tight_layout()
        try:
            from my_gui.fig_theme import apply_fig_dark_theme
            apply_fig_dark_theme(self._fig, self._style.bg())
        except Exception:
            pass
        self._canvas.draw()

    def _pick_rois(self):
        """Checkbox dialog to choose which ROIs the box plots include."""
        from PyQt6.QtWidgets import (QScrollArea, QDialogButtonBox)
        d = QDialog(self)
        d.setWindowTitle("Select ROIs")
        d.resize(280, 420)
        dv = QVBoxLayout(d)
        dv.addWidget(QLabel("Show box plots for these ROIs:"))
        sc = QScrollArea(); sc.setWidgetResizable(True)
        inner = QWidget(); il = QVBoxLayout(inner)
        chks = []
        for i, r in enumerate(self._rois):
            cb = QCheckBox(r.name); cb.setChecked(i in self._selected)
            il.addWidget(cb); chks.append(cb)
        il.addStretch()
        sc.setWidget(inner); dv.addWidget(sc, 1)

        anrow = QHBoxLayout()
        b_all = QPushButton("All")
        b_all.clicked.connect(lambda: [c.setChecked(True) for c in chks])
        b_none = QPushButton("None")
        b_none.clicked.connect(lambda: [c.setChecked(False) for c in chks])
        anrow.addWidget(b_all); anrow.addWidget(b_none); anrow.addStretch()
        dv.addLayout(anrow)

        bb = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok |
                              QDialogButtonBox.StandardButton.Cancel)
        bb.accepted.connect(d.accept); bb.rejected.connect(d.reject)
        dv.addWidget(bb)
        if d.exec() == QDialog.DialogCode.Accepted:
            sel = [i for i, c in enumerate(chks) if c.isChecked()]
            if not sel:
                QMessageBox.information(self, "No ROIs",
                                        "Select at least one ROI.")
                return
            self._selected = sel
            self._redraw()

    def _save(self):
        from my_gui.fig_export import save_figure, FIG_EXPORT_FILTER
        p, _ = QFileDialog.getSaveFileName(
            self, "Save box plots", "roi_boxplots", FIG_EXPORT_FILTER)
        if p:
            save_figure(self._fig, p, dpi=300, facecolor="white")

    def _run_stats(self):
        """Assign the (selected) ROIs to groups, then compare the groups within
        each map — choosing the test by group count + normality — and star the
        significant columns on the plots."""
        rois_idx = self._selected or list(range(len(self._rois)))

        # ── Group-assignment dialog ───────────────────────────────────────────
        d = QDialog(self)
        d.setWindowTitle("Group comparison")
        d.resize(380, 470)
        dv = QVBoxLayout(d)
        dv.addWidget(QLabel(
            "Assign each ROI to a group (type a label, e.g. 'Control' / "
            "'Treatment').\nROIs sharing a label form one group; blank = exclude.\n\n"
            "•  2 groups  → t-test (normal) or Mann-Whitney U (non-parametric)\n"
            "•  3+ groups → one-way ANOVA (normal) or Kruskal-Wallis"))
        tbl = QTableWidget(len(rois_idx), 2)
        tbl.setHorizontalHeaderLabels(["ROI", "Group"])
        tbl.verticalHeader().setVisible(False)
        tbl.horizontalHeader().setStretchLastSection(True)
        for r, i in enumerate(rois_idx):
            it = QTableWidgetItem(self._rois[i].name)
            it.setFlags(Qt.ItemFlag.ItemIsEnabled)
            tbl.setItem(r, 0, it)
            tbl.setItem(r, 1, QTableWidgetItem(self._groups.get(i, "")))
        tbl.resizeColumnsToContents()
        dv.addWidget(tbl, 1)
        brow = QHBoxLayout()
        b_ok = QPushButton("Run"); b_cancel = QPushButton("Cancel")
        brow.addStretch(); brow.addWidget(b_cancel); brow.addWidget(b_ok)
        dv.addLayout(brow)
        b_cancel.clicked.connect(d.reject); b_ok.clicked.connect(d.accept)
        if not d.exec():
            return

        # ── Collect labels ────────────────────────────────────────────────────
        self._groups = {}
        for r, i in enumerate(rois_idx):
            it = tbl.item(r, 1)
            lbl = it.text().strip() if it else ""
            if lbl:
                self._groups[i] = lbl
        labels = sorted(set(self._groups.values()))
        if len(labels) < 2:
            QMessageBox.information(self, "Need groups",
                                    "Assign at least two different group labels.")
            return

        # ── Compare the groups per map column (pooled pixels per group) ───────
        self._stats = {}
        lines = []
        for (pool, meth) in self._cols:
            arr = self._idx.get((pool, meth))
            if arr is None:
                continue
            group_data = []
            for lbl in labels:
                vals = [_roi_pixels(self._rois[i], arr)
                        for i, gl in self._groups.items()
                        if gl == lbl and _roi_pixels(self._rois[i], arr).size]
                group_data.append(np.concatenate(vals) if vals
                                  else np.asarray([], float))
            res = _group_compare(group_data)
            if res is not None:
                self._stats[(pool, meth)] = res
                _base = meth if not pool else f"{pool}: {meth}"
                lines.append(f"{_base}:  {res['test']}   p = {res['p']:.4g}   "
                             f"{_sig_stars(res['p'])}")

        self._redraw()
        if lines:
            QMessageBox.information(
                self, "Group comparison",
                f"Groups ({len(labels)}): {', '.join(labels)}\n\n"
                + "\n".join(lines)
                + "\n\n*  p < 0.05      **  p < 0.005\n"
                "(groups compared on pooled ROI pixels)")
        else:
            QMessageBox.information(
                self, "Group comparison",
                "No map column had enough data in ≥2 groups to compare.")


# ─────────────────────────────────────────────────────────────────────────────
# Convenience entry-point
# ─────────────────────────────────────────────────────────────────────────────

def show_roi_table(parent, rois: list,
                   map_items: list[tuple[str, np.ndarray | None]],
                   title: str = "ROI Statistics"):
    """
    Build and show a RoiTableDialog (non-modal).

    Automatically filters out 'Phantom_outline' ROIs and None/non-array maps.
    """
    clean_rois = [r for r in rois if r.name != "Phantom_outline"]
    valid_maps = [
        (lbl, arr) for lbl, arr in map_items
        if arr is not None and isinstance(arr, np.ndarray) and arr.ndim >= 2
    ]

    if not clean_rois:
        QMessageBox.information(parent, "No ROIs",
                                "Draw ROIs in the ROI Manager tab first.")
        return
    if not valid_maps:
        QMessageBox.information(parent, "No Maps",
                                "Run the analysis first to generate maps.")
        return

    dlg = RoiTableDialog(parent, clean_rois, valid_maps, title=title)
    dlg.show()
